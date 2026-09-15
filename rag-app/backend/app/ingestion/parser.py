from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pymupdf4llm
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from striprtf.striprtf import rtf_to_text


@dataclass
class ParsedPage:
    page_number: int  # 1-indexed
    markdown: str


def parse_document(path: str) -> list[ParsedPage]:
    """Dispatches on file extension to the right format-specific parser.
    Every parser below produces the same per-page Markdown shape regardless
    of source format -- the chunker and everything downstream only ever see
    Markdown text, so adding a new format means adding one function + one
    dispatch entry here, nothing else in the pipeline changes.
    """
    suffix = Path(path).suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"unsupported file type: {suffix!r} (supported: {', '.join(sorted(_PARSERS))})")
    return parser(path)


def _rows_to_markdown_table(rows: list[list[str]]) -> str:
    """First row is treated as the header. Produces the exact pipe-table
    syntax `chunker.py`'s `_TABLE_LINE` regex expects, so table rows get
    isolated as atomic chunks the same way for every source format.
    """
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    header = rows[0] + [""] * (width - len(rows[0]))
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join([" --- "] * width) + "|"]
    for row in rows[1:]:
        padded = list(row) + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded[:width]) + " |")
    return "\n".join(lines)


# --- PDF -------------------------------------------------------------


def _parse_pdf(path: str) -> list[ParsedPage]:
    """Per-page Markdown, preserving headings and rendering tables as
    Markdown tables so the chunker can reason about structure instead of
    raw, layout-stripped text.
    """
    pages = pymupdf4llm.to_markdown(path, page_chunks=True)
    return [ParsedPage(page_number=i + 1, markdown=page["text"]) for i, page in enumerate(pages)]


# --- Word (.docx) ------------------------------------------------------

_DOCX_HEADING_PREFIX = {
    "Title": "#", "Heading 1": "#", "Heading 2": "##", "Heading 3": "###", "Heading 4": "####",
}


def _iter_docx_blocks(document: Document):
    """python-docx exposes `.paragraphs` and `.tables` as two separate,
    order-losing lists -- walking the underlying XML body children directly
    is the documented way to get paragraphs and tables interleaved in their
    real document order.
    """
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def _parse_docx(path: str) -> list[ParsedPage]:
    """No native page concept in a .docx (page breaks are a rendering-time
    detail, not a reliable structural boundary) -- the whole document is one
    logical page. Section headings still carry through as Markdown headers,
    which is what the chunker actually keys off for chunk metadata.
    """
    document = Document(path)
    lines: list[str] = []
    for block in _iter_docx_blocks(document):
        if isinstance(block, DocxParagraph):
            text = block.text.strip()
            if not text:
                continue
            style_name = block.style.name if block.style else ""
            prefix = _DOCX_HEADING_PREFIX.get(style_name)
            lines.append(f"{prefix} {text}" if prefix else text)
        else:
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in block.rows]
            table_md = _rows_to_markdown_table(rows)
            if table_md:
                lines.append(table_md)
        lines.append("")
    return [ParsedPage(page_number=1, markdown="\n".join(lines))]


# --- PowerPoint (.pptx) ------------------------------------------------


def _pptx_table_to_markdown(table) -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    return _rows_to_markdown_table(rows)


_PPTX_TITLE_PLACEHOLDER_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _parse_pptx(path: str) -> list[ParsedPage]:
    """Unlike Word, a slide deck DOES have a natural page unit -- one
    ParsedPage per slide, matching how PDF pages map 1:1 to pages.
    """
    presentation = Presentation(path)
    pages: list[ParsedPage] = []
    for i, slide in enumerate(presentation.slides):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                # `shape is slide.shapes.title` never matches -- python-pptx
                # hands back a fresh wrapper object each access, even for the
                # same underlying shape -- so the title is identified by its
                # placeholder TYPE instead, which is a stable property.
                is_title = shape.is_placeholder and shape.placeholder_format.type in _PPTX_TITLE_PLACEHOLDER_TYPES
                lines.append(f"# {text}" if is_title else text)
            elif shape.has_table:
                table_md = _pptx_table_to_markdown(shape.table)
                if table_md:
                    lines.append(table_md)
        pages.append(ParsedPage(page_number=i + 1, markdown="\n\n".join(lines)))
    return pages


# --- Plain text-like formats (Markdown, .txt) ---------------------------


def _parse_plain_text(path: str) -> list[ParsedPage]:
    """Markdown files are already exactly the format the chunker expects;
    plain .txt has no heading/table syntax to preserve, so it's read as-is
    too -- the chunker still splits it sensibly by paragraph/token size,
    it just won't have section headings to key chunk metadata off of.
    """
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return [ParsedPage(page_number=1, markdown=text)]


# --- Spreadsheets (.xlsx, .csv) -----------------------------------------


def _parse_xlsx(path: str) -> list[ParsedPage]:
    """One page per sheet -- another format with a natural page-like unit."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    pages: list[ParsedPage] = []
    for i, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        rows = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(cell.strip() for cell in row)]
        if not rows:
            continue
        pages.append(ParsedPage(page_number=i + 1, markdown=f"## {sheet_name}\n\n{_rows_to_markdown_table(rows)}"))
    return pages


def _parse_csv(path: str) -> list[ParsedPage]:
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        rows = [row for row in csv.reader(f) if any(cell.strip() for cell in row)]
    if not rows:
        return []
    return [ParsedPage(page_number=1, markdown=_rows_to_markdown_table(rows))]


# --- Rich Text Format (.rtf) ---------------------------------------------


def _parse_rtf(path: str) -> list[ParsedPage]:
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    return [ParsedPage(page_number=1, markdown=rtf_to_text(raw))]


# --- HTML (.html, .htm) ---------------------------------------------------

_HTML_HEADING_PREFIX = {f"h{level}": "#" * level for level in range(1, 7)}


def _parse_html(path: str) -> list[ParsedPage]:
    """Simplification, not a full HTML->Markdown converter: walks top-level
    heading/paragraph/list-item/table tags in document order. Nested blocks
    (e.g. a <table> inside a <td>) can be picked up more than once -- an
    accepted simplification given how rarely deeply-nested HTML shows up in
    an HR knowledge base versus a straightforward exported article/page.
    """
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    lines: list[str] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"]):
        if element.name in _HTML_HEADING_PREFIX:
            text = element.get_text(strip=True)
            if text:
                lines.append(f"{_HTML_HEADING_PREFIX[element.name]} {text}")
        elif element.name == "table":
            rows = [[cell.get_text(strip=True) for cell in row.find_all(["td", "th"])] for row in element.find_all("tr")]
            rows = [row for row in rows if row]
            table_md = _rows_to_markdown_table(rows)
            if table_md:
                lines.append(table_md)
        else:
            text = element.get_text(strip=True)
            if text:
                lines.append(text)
        lines.append("")
    return [ParsedPage(page_number=1, markdown="\n".join(lines))]


_PARSERS: dict[str, Callable[[str], list[ParsedPage]]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
    ".md": _parse_plain_text,
    ".markdown": _parse_plain_text,
    ".txt": _parse_plain_text,
    ".xlsx": _parse_xlsx,
    ".csv": _parse_csv,
    ".rtf": _parse_rtf,
    ".html": _parse_html,
    ".htm": _parse_html,
}

SUPPORTED_EXTENSIONS = frozenset(_PARSERS)
